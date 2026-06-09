import io
import json
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, field_validator
from sqlalchemy.orm import Session, selectinload

from ..auth import get_current_admin, get_current_user
from ..database import get_db
from ..models import AppSettings, Eintrag, Tier, User, eintrag_tiere

router = APIRouter(prefix="/api/reittagebuch", tags=["reittagebuch"])


# ── Schemas ──────────────────────────────────────────────────────

class ZeitSlot(BaseModel):
    von: str
    bis: str


class TierCreate(BaseModel):
    name: str
    typ: str
    emoji: str = "🐴"


class TierUpdate(BaseModel):
    name: Optional[str] = None
    typ: Optional[str] = None
    emoji: Optional[str] = None
    aktiv: Optional[bool] = None


class TierOut(BaseModel):
    id: int
    name: str
    typ: str
    emoji: str
    aktiv: bool
    created_at: datetime
    model_config = {"from_attributes": True}


class EintragUserOut(BaseModel):
    display_name: str
    model_config = {"from_attributes": True}


class EintragCreate(BaseModel):
    datum: date
    zeiten: Optional[list[ZeitSlot]] = None
    aktivitaet: str
    besonderheiten: Optional[str] = None
    anpassungen: Optional[str] = None
    anzahl_kinder: int = 0
    anzahl_jugendliche: int = 0
    tier_ids: list[int] = []


class EintragOut(BaseModel):
    id: int
    datum: date
    zeiten: Optional[list[ZeitSlot]] = None
    aktivitaet: str
    besonderheiten: Optional[str]
    anpassungen: Optional[str]
    anzahl_kinder: int
    anzahl_jugendliche: int
    tiere: list[TierOut]
    user: Optional[EintragUserOut] = None
    created_at: datetime
    updated_at: datetime

    @field_validator('zeiten', mode='before')
    @classmethod
    def parse_zeiten(cls, v):
        if isinstance(v, str):
            return json.loads(v) if v else None
        return v

    model_config = {"from_attributes": True}


class WochenDurchschnitt(BaseModel):
    kinder: float
    jugendliche: float
    gesamt: float


class TierEinsatz(BaseModel):
    tier_id: int
    name: str
    emoji: str
    anzahl: int


class WocheVerlauf(BaseModel):
    kw: str
    kinder: int
    jugendliche: int


class StatsOut(BaseModel):
    gesamt_einheiten: int
    gesamt_kinder: int
    gesamt_jugendliche: int
    wochen_durchschnitt: WochenDurchschnitt
    tiere_einsaetze: list[TierEinsatz]
    wochen_verlauf: list[WocheVerlauf]


class TierTypenUpdate(BaseModel):
    typen: list[str]


# ── Tier-Typen ───────────────────────────────────────────────────

TIER_TYPEN_KEY = "tier_typen"
TIER_TYPEN_DEFAULT = ["Pferd", "Pony", "Esel", "Maultier"]


def _get_tiertypen(db: Session) -> list[str]:
    row = db.query(AppSettings).filter(AppSettings.key == TIER_TYPEN_KEY).first()
    return json.loads(row.value) if row else TIER_TYPEN_DEFAULT


@router.get("/tiertypen", response_model=list[str])
def get_tiertypen(
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    return _get_tiertypen(db)


@router.put("/tiertypen", response_model=list[str])
def update_tiertypen(
    body: TierTypenUpdate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    cleaned = list(dict.fromkeys(t.strip() for t in body.typen if t.strip()))
    if not cleaned:
        raise HTTPException(status_code=400, detail="Mindestens ein Tiertyp erforderlich")
    serialized = json.dumps(cleaned)
    row = db.query(AppSettings).filter(AppSettings.key == TIER_TYPEN_KEY).first()
    if row is None:
        row = AppSettings(key=TIER_TYPEN_KEY, value=serialized, updated_by=current_user.id)
        db.add(row)
    else:
        row.value = serialized
        row.updated_by = current_user.id
        row.updated_at = datetime.utcnow()
    db.commit()
    return cleaned


# ── Helpers ──────────────────────────────────────────────────────

def _resolve_tiere(tier_ids: list[int], db: Session) -> list[Tier]:
    if not tier_ids:
        return []
    unique_ids = list(set(tier_ids))
    tiere = db.query(Tier).filter(Tier.id.in_(unique_ids)).all()
    if len(tiere) != len(unique_ids):
        raise HTTPException(status_code=400, detail="Ein oder mehrere Tiere nicht gefunden")
    return tiere


def _eintraege_query(
    von: Optional[date],
    bis: Optional[date],
    tier_id: Optional[int],
    db: Session,
):
    q = (
        db.query(Eintrag)
        .options(selectinload(Eintrag.tiere), selectinload(Eintrag.user))
    )
    if von:
        q = q.filter(Eintrag.datum >= von)
    if bis:
        q = q.filter(Eintrag.datum <= bis)
    if tier_id:
        q = q.filter(Eintrag.tiere.any(Tier.id == tier_id))
    return q


def _reload(eintrag_id: int, db: Session) -> Eintrag:
    return (
        db.query(Eintrag)
        .options(selectinload(Eintrag.tiere), selectinload(Eintrag.user))
        .filter(Eintrag.id == eintrag_id)
        .first()
    )


# ── Tiere ────────────────────────────────────────────────────────

@router.get("/tiere", response_model=list[TierOut])
def list_tiere(
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    return (
        db.query(Tier)
        .order_by(Tier.aktiv.desc(), Tier.name)
        .all()
    )


@router.post("/tiere", response_model=TierOut, status_code=status.HTTP_201_CREATED)
def create_tier(
    body: TierCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    tier = Tier(user_id=current_user.id, **body.model_dump())
    db.add(tier)
    db.commit()
    db.refresh(tier)
    return tier


@router.put("/tiere/{tier_id}", response_model=TierOut)
def update_tier(
    tier_id: int,
    body: TierUpdate,
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    tier = db.query(Tier).filter(Tier.id == tier_id).first()
    if not tier:
        raise HTTPException(status_code=404, detail="Tier nicht gefunden")
    for field, value in body.model_dump(exclude_none=True).items():
        setattr(tier, field, value)
    db.commit()
    db.refresh(tier)
    return tier


@router.delete("/tiere/{tier_id}", status_code=status.HTTP_204_NO_CONTENT)
def deactivate_tier(
    tier_id: int,
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    tier = db.query(Tier).filter(Tier.id == tier_id).first()
    if not tier:
        raise HTTPException(status_code=404, detail="Tier nicht gefunden")
    tier.aktiv = False
    db.commit()


# ── Einträge ─────────────────────────────────────────────────────

@router.get("/eintraege", response_model=list[EintragOut])
def list_eintraege(
    von: Optional[date] = None,
    bis: Optional[date] = None,
    tier_id: Optional[int] = None,
    limit: int = Query(50, ge=1, le=500),
    offset: int = Query(0, ge=0),
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    return (
        _eintraege_query(von, bis, tier_id, db)
        .order_by(Eintrag.datum.desc(), Eintrag.created_at.desc())
        .limit(limit)
        .offset(offset)
        .all()
    )


@router.get("/eintraege/{eintrag_id}", response_model=EintragOut)
def get_eintrag(
    eintrag_id: int,
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    eintrag = _reload(eintrag_id, db)
    if not eintrag:
        raise HTTPException(status_code=404, detail="Eintrag nicht gefunden")
    return eintrag


@router.post("/eintraege", response_model=EintragOut, status_code=status.HTTP_201_CREATED)
def create_eintrag(
    body: EintragCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    tiere = _resolve_tiere(body.tier_ids, db)
    eintrag = Eintrag(
        user_id=current_user.id,
        datum=body.datum,
        zeiten=json.dumps([z.model_dump() for z in body.zeiten]) if body.zeiten else None,
        aktivitaet=body.aktivitaet,
        besonderheiten=body.besonderheiten,
        anpassungen=body.anpassungen,
        anzahl_kinder=body.anzahl_kinder,
        anzahl_jugendliche=body.anzahl_jugendliche,
        tiere=tiere,
    )
    db.add(eintrag)
    db.commit()
    return _reload(eintrag.id, db)


@router.put("/eintraege/{eintrag_id}", response_model=EintragOut)
def update_eintrag(
    eintrag_id: int,
    body: EintragCreate,
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    eintrag = (
        db.query(Eintrag)
        .options(selectinload(Eintrag.tiere))
        .filter(Eintrag.id == eintrag_id)
        .first()
    )
    if not eintrag:
        raise HTTPException(status_code=404, detail="Eintrag nicht gefunden")

    tiere = _resolve_tiere(body.tier_ids, db)

    eintrag.datum = body.datum
    eintrag.zeiten = json.dumps([z.model_dump() for z in body.zeiten]) if body.zeiten else None
    eintrag.aktivitaet = body.aktivitaet
    eintrag.besonderheiten = body.besonderheiten
    eintrag.anpassungen = body.anpassungen
    eintrag.anzahl_kinder = body.anzahl_kinder
    eintrag.anzahl_jugendliche = body.anzahl_jugendliche
    eintrag.tiere = tiere
    eintrag.updated_at = datetime.utcnow()

    db.commit()
    return _reload(eintrag_id, db)


@router.delete("/eintraege/{eintrag_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_eintrag(
    eintrag_id: int,
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    eintrag = db.query(Eintrag).filter(Eintrag.id == eintrag_id).first()
    if not eintrag:
        raise HTTPException(status_code=404, detail="Eintrag nicht gefunden")
    db.delete(eintrag)
    db.commit()


# ── Auswertungen ─────────────────────────────────────────────────

@router.get("/stats", response_model=StatsOut)
def get_stats(
    von: Optional[date] = None,
    bis: Optional[date] = None,
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    eintraege = _eintraege_query(von, bis, None, db).all()

    gesamt_kinder = sum(e.anzahl_kinder for e in eintraege)
    gesamt_jugendliche = sum(e.anzahl_jugendliche for e in eintraege)
    gesamt_einheiten = len(eintraege)

    if eintraege:
        d_min = von or min(e.datum for e in eintraege)
        d_max = bis or max(e.datum for e in eintraege)
        num_weeks = max(1.0, (d_max - d_min).days / 7.0)
    else:
        num_weeks = 1.0

    wochen_durchschnitt = WochenDurchschnitt(
        kinder=round(gesamt_kinder / num_weeks, 1),
        jugendliche=round(gesamt_jugendliche / num_weeks, 1),
        gesamt=round(gesamt_einheiten / num_weeks, 1),
    )

    tier_counts: dict[int, dict] = {}
    for e in eintraege:
        for t in e.tiere:
            if t.id not in tier_counts:
                tier_counts[t.id] = {"tier_id": t.id, "name": t.name, "emoji": t.emoji, "anzahl": 0}
            tier_counts[t.id]["anzahl"] += 1
    tiere_einsaetze = sorted(tier_counts.values(), key=lambda x: x["anzahl"], reverse=True)

    wochen: dict[tuple, dict] = {}
    for e in eintraege:
        year, kw, _ = e.datum.isocalendar()
        key = (year, kw)
        if key not in wochen:
            wochen[key] = {"kw": f"KW {kw:02d}", "kinder": 0, "jugendliche": 0}
        wochen[key]["kinder"] += e.anzahl_kinder
        wochen[key]["jugendliche"] += e.anzahl_jugendliche
    wochen_verlauf = [wochen[k] for k in sorted(wochen)]

    return StatsOut(
        gesamt_einheiten=gesamt_einheiten,
        gesamt_kinder=gesamt_kinder,
        gesamt_jugendliche=gesamt_jugendliche,
        wochen_durchschnitt=wochen_durchschnitt,
        tiere_einsaetze=[TierEinsatz(**t) for t in tiere_einsaetze],
        wochen_verlauf=[WocheVerlauf(**w) for w in wochen_verlauf],
    )


# ── Excel-Export ─────────────────────────────────────────────────

@router.get("/export")
def export_xlsx(
    von: Optional[date] = None,
    bis: Optional[date] = None,
    _: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    eintraege = (
        _eintraege_query(von, bis, None, db)
        .order_by(Eintrag.datum.asc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoftagebuch"

    headers = ["Datum", "Zeiten", "Tiere", "Aktivität", "Kinder", "Jugendliche", "Besonderheiten", "Anpassungen", "Erstellt von"]
    ws.append(headers)

    header_fill = PatternFill(start_color="5B7C5E", end_color="5B7C5E", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=False)

    for e in eintraege:
        tiere_str = ", ".join(f"{t.emoji} {t.name}" for t in e.tiere) if e.tiere else "—"
        zeiten_raw = json.loads(e.zeiten) if e.zeiten else []
        zeiten_str = ", ".join(f"{z['von']}–{z['bis']}" for z in zeiten_raw) if zeiten_raw else ""
        ws.append([
            e.datum,
            zeiten_str,
            tiere_str,
            e.aktivitaet,
            e.anzahl_kinder,
            e.anzahl_jugendliche,
            e.besonderheiten or "",
            e.anpassungen or "",
            e.user.display_name if e.user else "",
        ])

    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "DD.MM.YYYY"

    for i, width in enumerate([13, 18, 22, 45, 9, 14, 38, 38, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    von_str = von.isoformat() if von else "alle"
    bis_str = bis.isoformat() if bis else "alle"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="hoftagebuch_{von_str}_{bis_str}.xlsx"'},
    )


# ── DB-Backup ────────────────────────────────────────────────────

@router.get("/backup")
def download_backup(_: User = Depends(get_current_admin)):
    from pathlib import Path
    db_path = Path(__file__).resolve().parent.parent.parent / "data" / "toolbox.db"
    if not db_path.exists():
        raise HTTPException(status_code=404, detail="Datenbank-Datei nicht gefunden")
    return FileResponse(
        path=str(db_path),
        media_type="application/octet-stream",
        filename="toolbox.db",
        headers={"Content-Disposition": 'attachment; filename="toolbox.db"'},
    )
